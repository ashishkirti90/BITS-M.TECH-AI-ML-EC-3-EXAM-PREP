from pathlib import Path
import logging, re
from pypdf import PdfReader
from docx import Document

logging.getLogger("pypdf").setLevel(logging.ERROR)
root=Path(r"C:\Users\Admin\Downloads\Sem 1")
src=root/"LATEST QUESTION PAPERS"
out=root/"exam_analysis"/"latest_extracted"
out.mkdir(parents=True,exist_ok=True)

for p in src.rglob('*'):
    if not p.is_file() or 'EC3' not in p.name: continue
    parts=[f'SOURCE: {p.relative_to(root)}']
    try:
        if p.suffix.lower()=='.pdf':
            r=PdfReader(str(p)); parts.append(f'PAGES: {len(r.pages)}')
            for i,page in enumerate(r.pages,1): parts.append(f'\n===== PAGE {i} =====\n{page.extract_text() or ""}')
        elif p.suffix.lower()=='.docx':
            d=Document(str(p))
            parts += [x.text for x in d.paragraphs if x.text.strip()]
            for ti,t in enumerate(d.tables,1):
                parts.append(f'\n===== TABLE {ti} =====')
                parts += ['\t'.join(c.text.replace('\n',' | ') for c in row.cells) for row in t.rows]
        else:
            parts.append('[LEGACY DOC REQUIRES CONVERSION]')
        target=out/(re.sub(r'[^A-Za-z0-9._-]+','_',str(p.relative_to(root)))+'.txt')
        target.write_text('\n'.join(parts),encoding='utf-8')
        print(p.name,'->',target.name,len('\n'.join(parts)))
    except Exception as e: print('ERROR',p,e)
