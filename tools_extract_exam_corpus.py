from pathlib import Path
import re
from pypdf import PdfReader
from docx import Document

ROOT = Path(r"C:\Users\Admin\Downloads\Sem 1")
OUT = ROOT / "exam_analysis" / "extracted"
OUT.mkdir(parents=True, exist_ok=True)

def wanted(p: Path) -> bool:
    s = str(p).lower()
    name = p.name.lower()
    if "course handout" in name:
        return True
    if "watermark" in name or "watermarked" in name:
        return False
    if "question papers" not in s and "previous question papers" not in s:
        return False
    end_tokens = ("endsem", "end-sem", "comprehensive", "compre", "ec3", "past papers", "sample")
    return any(t in name for t in end_tokens)

def extract_pdf(p: Path) -> str:
    reader = PdfReader(str(p))
    parts = []
    for i, page in enumerate(reader.pages, 1):
        try:
            text = page.extract_text() or ""
        except Exception as e:
            text = f"[EXTRACTION ERROR: {e}]"
        parts.append(f"\n===== PAGE {i} =====\n{text}")
    return "".join(parts)

def extract_docx(p: Path) -> str:
    doc = Document(str(p))
    parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    for ti, table in enumerate(doc.tables, 1):
        parts.append(f"\n===== TABLE {ti} =====")
        for row in table.rows:
            parts.append("\t".join(cell.text.replace("\n", " | ") for cell in row.cells))
    return "\n".join(parts)

manifest = []
for p in sorted(ROOT.rglob("*")):
    if not p.is_file() or p.suffix.lower() not in {".pdf", ".docx"} or not wanted(p):
        continue
    rel = p.relative_to(ROOT)
    safe = re.sub(r"[^A-Za-z0-9._-]+", "_", str(rel)) + ".txt"
    out = OUT / safe
    if out.exists() and out.stat().st_size > 100:
        manifest.append(f"SKIP\t{rel}\t{out.stat().st_size}\t{out.name}")
        continue
    try:
        text = extract_pdf(p) if p.suffix.lower() == ".pdf" else extract_docx(p)
        out.write_text(f"SOURCE: {rel}\n{text}", encoding="utf-8")
        manifest.append(f"OK\t{rel}\t{len(text)}\t{out.name}")
    except Exception as e:
        manifest.append(f"ERROR\t{rel}\t{e}")
(OUT.parent / "manifest.tsv").write_text("\n".join(manifest), encoding="utf-8")
print(f"Processed {len(manifest)} files; output={OUT}")
